#!/usr/bin/env python3
# ============================================================
# mark_phase4_complete.py — تعليم المرحلة 04 منجزة في خطة أوج
# (نفس نمط الجراحة النصية لمراحل 0-3)
#   1. بطاقة المتابعة: [x] منجزة + التواريخ + المنفّذ
#   2. 8 مهام [x]
#   3. صف اللوحة 05: أخضر (2E7D32 + تعبئة E8F5E9 + غامق) + تواريخ + ملاحظة أدلة
# ============================================================
from docx import Document
from docx.shared import RGBColor
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
PATH = 'docs/Awj_Development_and_Launch_Plan.docx'
DATE_AR = '11 سبتمبر 2026'
DATE_ISO = '2026-09-11'

doc = Document(PATH)
changes = []

# ── 1. بطاقة المرحلة الرابعة (TABLE 6) ──
card = doc.tables[6]
p0 = card.rows[0].cells[0].paragraphs[0]
r0 = p0.runs[0]
assert 'الرابعة' in r0.text and '[ ] منجزة' in r0.text, f'card p0 unexpected: {r0.text!r}'
r0.text = r0.text.replace('[ ] منجزة', '[x] منجزة')
changes.append('card: [x] منجزة')

p1 = card.rows[0].cells[0].paragraphs[1]
r1 = p1.runs[0]
assert 'تاريخ البدء' in r1.text
r1.text = (
    f'تاريخ البدء: {DATE_AR}   تاريخ الإنجاز: {DATE_AR}   المنفّذ: Super Z'
)
changes.append('card: dates + Super Z')

# ── 2. المهام الثماني (فقرات Checklist بعد عنوان Tasks الخاص بالمرحلة 04) ──
in_phase4 = False
marked = 0
for p in doc.paragraphs:
    t = p.text.strip()
    if 'المرحلة الرابعة' in t and 'Plans' in t:
        in_phase4 = True
        continue
    if in_phase4 and 'المرحلة الخامسة' in t:
        break
    if in_phase4 and '[ ]' in t:
        # أول run يبدأ بالصندوق (بنية الفقرات: run واحد غالبًا)
        done = False
        for run in p.runs:
            if '[ ]' in run.text:
                run.text = run.text.replace('[ ]', '[x]', 1)
                done = True
                break
        if not done:
            # fallback: دمج كل الـ runs ثم إعادة الكتابة في الأول
            full = ''.join(run.text for run in p.runs).replace('[ ]', '[x]', 1)
            p.runs[0].text = full
            for run in p.runs[1:]:
                run.text = ''
        marked += 1
assert marked == 8, f'expected 8 tasks, marked {marked}'
changes.append(f'tasks: 8 × [x]')

# ── 3. اللوحة: صف 05 (TABLE 0, row 5) ──
board = doc.tables[0]
row = board.rows[5]
cells = row.cells
assert cells[0].text.strip() == '05', f'row mismatch: {cells[0].text!r}'
assert cells[2].text.strip() == 'لم تبدأ', f'status mismatch: {cells[2].text!r}'

# الحالة → منجزة (غامق + أخضر 2E7D32 + تعبئة E8F5E9)
status_cell = cells[2]
status_run = status_cell.paragraphs[0].runs[0]
status_run.text = 'منجزة'
status_run.bold = True
status_run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
tcPr = status_cell._tc.get_or_add_tcPr()
# أزل أي shd سابق
for shd in tcPr.findall(f'{{{W}}}shd'):
    tcPr.remove(shd)
shd = etree.SubElement(tcPr, f'{{{W}}}shd')
shd.set(f'{{{W}}}val', 'clear')
shd.set(f'{{{W}}}fill', 'E8F5E9')
changes.append('board row 05: منجزة green+bold+fill')

# التواريخ
for ci, val in [(3, DATE_ISO), (4, DATE_ISO)]:
    cell = cells[ci]
    if cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = val
    else:
        cell.paragraphs[0].add_run(val)
changes.append('board row 05: dates')

# ملاحظة الأدلة
note_cell = cells[5]
note = (
    'منجز كاملًا: Monetization Core — مصدر واحد للحدود (plan_entitlements) + '
    'enforcement ذرّي في القاعدة (consume_usage) + عدادات server-side + '
    'عرض الاستخدام وUpgrade Prompts + دفع يدوي بمراجعة أدمن — الأدلة: '
    'docs/phase-4/MONETIZATION.md'
)
if note_cell.paragraphs[0].runs:
    note_cell.paragraphs[0].runs[0].text = note
else:
    note_cell.paragraphs[0].add_run(note)
changes.append('board row 05: evidence note')

doc.save(PATH)
print('✅ marked phase 4 complete:')
for c in changes:
    print('  •', c)
