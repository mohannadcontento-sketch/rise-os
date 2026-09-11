#!/usr/bin/env python3
# ============================================================
# mark_phase5_complete.py — تعليم المرحلة 05 منجزة في خطة أوج
# (نفس نمط الجراحة النصية لمراحل 0-4)
#   1. بطاقة المتابعة (TABLE 8): [x] منجزة + التواريخ + المنفّذ
#   2. 15 بندًا [x] (8 مهام + 7 أنواع أحداث)
#   3. صف اللوحة 06: أخضر (2E7D32 + تعبئة E8F5E9 + غامق) + تواريخ + ملاحظة أدلة
# ============================================================
from docx import Document
from docx.shared import RGBColor
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
PATH = 'docs/Awj_Development_and_Launch_Plan.docx'
DATE_AR = '12 سبتمبر 2026'
DATE_ISO = '2026-09-12'

doc = Document(PATH)
changes = []

# ── 1. بطاقة المرحلة الخامسة (TABLE 8) ──
card = doc.tables[8]
p0 = card.rows[0].cells[0].paragraphs[0]
r0 = p0.runs[0]
assert 'الخامسة' in r0.text and '[ ] منجزة' in r0.text, f'card p0 unexpected: {r0.text!r}'
r0.text = r0.text.replace('[ ] منجزة', '[x] منجزة')
changes.append('card: [x] منجزة')

p1 = card.rows[0].cells[0].paragraphs[1]
r1 = p1.runs[0]
assert 'تاريخ البدء' in r1.text
r1.text = (
    f'تاريخ البدء: {DATE_AR}   تاريخ الإنجاز: {DATE_AR}   المنفّذ: Super Z'
)
changes.append('card: dates + Super Z')

# ── 2. البنود (8 مهام + 7 أنواع أحداث) بين عنواني المرحلتين 05 و06 ──
in_phase5 = False
marked = 0
for p in doc.paragraphs:
    t = p.text.strip()
    if 'المرحلة الخامسة' in t and 'الإشعارات' in t:
        in_phase5 = True
        continue
    if in_phase5 and 'المرحلة السادسة' in t:
        break
    if in_phase5 and '[ ]' in t:
        done = False
        for run in p.runs:
            if '[ ]' in run.text:
                run.text = run.text.replace('[ ]', '[x]', 1)
                done = True
                break
        if not done:
            full = ''.join(run.text for run in p.runs).replace('[ ]', '[x]', 1)
            p.runs[0].text = full
            for run in p.runs[1:]:
                run.text = ''
        marked += 1
assert marked == 15, f'expected 15 items, marked {marked}'
changes.append(f'items: 15 × [x] (8 مهام + 7 أنواع)')

# ── 3. اللوحة: صف 06 (TABLE 0, row 6) ──
board = doc.tables[0]
row = board.rows[6]
cells = row.cells
assert cells[0].text.strip() == '06', f'row mismatch: {cells[0].text!r}'
assert cells[2].text.strip() == 'لم تبدأ', f'status mismatch: {cells[2].text!r}'

# الحالة → منجزة (غامق + أخضر 2E7D32 + تعبئة E8F5E9)
status_cell = cells[2]
status_run = status_cell.paragraphs[0].runs[0]
status_run.text = 'منجزة'
status_run.bold = True
status_run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
tcPr = status_cell._tc.get_or_add_tcPr()
for shd in tcPr.findall(f'{{{W}}}shd'):
    tcPr.remove(shd)
shd = etree.SubElement(tcPr, f'{{{W}}}shd')
shd.set(f'{{{W}}}val', 'clear')
shd.set(f'{{{W}}}fill', 'E8F5E9')
changes.append('board row 06: منجزة green+bold+fill')

# التواريخ
for ci, val in [(3, DATE_ISO), (4, DATE_ISO)]:
    cell = cells[ci]
    if cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = val
    else:
        cell.paragraphs[0].add_run(val)
changes.append('board row 06: dates')

# ملاحظة الأدلة
note_cell = cells[5]
note = (
    'منجز كاملًا: مركز إشعارات موحد — notify_user (بوابة صلاحيات في القاعدة) + '
    'dedup (فهرس فريد) + فلاتر server-side + حذف كسول للمنتهي + إشعارات الحدود '
    'من consume_usage + تاب قالب إيميل الريست في لوحة الأدمن — الأدلة: '
    'docs/phase-5/NOTIFICATIONS.md'
)
if note_cell.paragraphs[0].runs:
    note_cell.paragraphs[0].runs[0].text = note
else:
    note_cell.paragraphs[0].add_run(note)
changes.append('board row 06: evidence note')

doc.save(PATH)
print('✅ marked phase 5 complete:')
for c in changes:
    print('  •', c)
